from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

wb = Workbook()
F = 'Arial'
blue = Font(name=F, color='0000FF')
black = Font(name=F)
bold = Font(name=F, bold=True)
green = Font(name=F, color='008000')
title = Font(name=F, bold=True, size=13)
yellow = PatternFill('solid', fgColor='FFFF00')
hdrfill = PatternFill('solid', fgColor='D9E1F2')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

# ---------- Comparison ----------
ws = wb.active
ws.title = 'Comparison'
# ---------- Inputs ----------
wi = wb.create_sheet('Inputs')
# ---------- Washers ----------
ww = wb.create_sheet('Washers')

# Inputs
wi['A1'] = 'Design inputs'; wi['A1'].font = title
inputs = [
 ('Preload target (kN)', 8.2, 'Design preload; Hoss guidance for TTS on small mills is 6.7–8.9 kN (1,500–2,000 lbf).'),
 ('Release stroke s_r (mm)', 2.5, 'Assumed until the release test with the new collet.'),
 ('Collet push-out force P (kN)', 0.5, 'Assumed; first estimate from the taper break test.'),
 ('Release deflection limit (× h0)', 0.65, 'Design choice, below the 0.75 h0 recommended maximum dynamic deflection (DIN 2093 guidance, per Surisa).'),
 ('Cylinder bore (mm)', 80, 'Selected cylinder.'),
 ('Design pressure (bar)', 5.5, 'Supply is 90 psi (6.2 bar); 5.5 bar allows for line drop and cut-in.'),
 ('Lever ratio i', 5, 'Proposed; adjustable 5–6:1.'),
 ('Lever efficiency η', 0.9, 'Assumed.'),
 ('Force available at drawbar (kN)', None, 'p × A × i × η'),
 ('Cap wall thickness (mm)', 6, 'Proposed.'),
 ('Washer OD guide clearance (mm)', 0.6, 'DIN 2093 guide tolerance for De > 31.5–50 mm (Surisa).'),
]
for i,(lab,val,note) in enumerate(inputs, start=2):
    wi.cell(i,1,lab).font = black
    c = wi.cell(i,2,val)
    c.font = blue
    wi.cell(i,3,note).font = Font(name=F, italic=True, color='595959')
wi['B10'] = '=PI()*B6^2/4*B7/10/1000*B8*B9'; wi['B10'].font = black; wi['B10'].number_format = '0.0'
for r in range(2,13):
    if r!=10: wi.cell(r,2).number_format = '0.00'
wi['B6'].number_format='0'; wi['B11'].number_format='0'
wi.column_dimensions['A'].width = 34; wi.column_dimensions['B'].width = 10; wi.column_dimensions['C'].width = 80

# Washers (DIN 2093 series A data)
ww['A1'] = 'DIN 2093 series A washer data'; ww['A1'].font = title
heads = ['Size','De (mm)','Di (mm)','t (mm)','h0 (mm)','l0 (mm)','F 0.25 h0 (kN)','F 0.5 h0 (kN)','F 0.75 h0 (kN)','F 1.0 h0 (kN)','Mass each (g)']
for j,h in enumerate(heads, start=1):
    c = ww.cell(2,j,h); c.font = bold; c.fill = hdrfill; c.border = box; c.alignment = Alignment(wrap_text=True)
rows = [
 ('A50 × 25.4 × 3', 50, 25.4, 3.0, 1.1, 4.1, 4.25, 8.21, 12.0, 15.6, 32.529),
 ('A40 × 20.4 × 2.25', 40, 20.4, 2.25, 0.9, 3.15, 2.34, 4.48, 6.50, 8.46, 15.26),
]
for i,row in enumerate(rows, start=3):
    for j,v in enumerate(row, start=1):
        c = ww.cell(i,j,v); c.border = box
        c.font = blue if j>1 else black
ww['A6'] = 'Source: DIN 2093 table as published by Fasten.it, https://www.fasten.it/en/norms/norm/din_2093 (forces converted N → kN; mass = kg per 1,000 pcs).'
ww['A6'].font = Font(name=F, italic=True, color='595959')
ww.column_dimensions['A'].width = 20
for col in 'BCDEFGHIJK': ww.column_dimensions[col].width = 11
ww.row_dimensions[2].height = 30

# Comparison
ws['A1'] = 'Belleville stack comparison: 16 × A50 single vs A40 in pairs'; ws['A1'].font = title
ws['A2'] = 'Blue = input you can change · black = formula · yellow = fill in from the quote'; ws['A2'].font = Font(name=F, italic=True, color='595959')
hdr = ['Parameter', 'A50 single', 'A40 in pairs']
for j,h in enumerate(hdr, start=1):
    c = ws.cell(4,j,h); c.font = bold; c.fill = hdrfill; c.border = box

def interp(x, r):
    W='Washers!'
    return (f'IF({x}<=0.25,{W}G{r}*{x}/0.25,IF({x}<=0.5,{W}G{r}+({x}-0.25)/0.25*({W}H{r}-{W}G{r}),'
            f'IF({x}<=0.75,{W}H{r}+({x}-0.5)/0.25*({W}I{r}-{W}H{r}),{W}I{r}+({x}-0.75)/0.25*({W}J{r}-{W}I{r}))))')
def inv(f, r):
    W='Washers!'
    return (f'IF({f}<={W}G{r},0.25*{f}/{W}G{r},IF({f}<={W}H{r},0.25+0.25*({f}-{W}G{r})/({W}H{r}-{W}G{r}),'
            f'IF({f}<={W}I{r},0.5+0.25*({f}-{W}H{r})/({W}I{r}-{W}H{r}),0.75+0.25*({f}-{W}I{r})/({W}J{r}-{W}I{r}))))')

labels = [
 ('Washer size', 'text'),
 ('Washers per group (parallel)', 'in'),
 ('Per-washer force at preload (kN)', '0.00'),
 ('Preload deflection (× h0)', '0.00'),
 ('Travel per group to limit (mm)', '0.000'),
 ('Minimum groups to stay within limit', '0'),
 ('Groups in series (chosen)', 'in'),
 ('Within deflection limit?', 'text'),
 ('Total washers', '0'),
 ('Release deflection (× h0)', '0.00'),
 ('Release force (kN)', '0.0'),
 ('Required at drawbar incl. P (kN)', '0.0'),
 ('Cylinder margin', '0%'),
 ('Group free height (mm)', '0.00'),
 ('Free stack length (mm)', '0.0'),
 ('Preloaded stack length (mm)', '0.0'),
 ('Stack-height guide, 3 × OD (mm)', '0'),
 ('Stack-height check', 'text'),
 ('Washer mass, total (g)', '0'),
 ('Cap washer bore (mm)', '0.0'),
 ('Cap OD, estimate (mm)', '0.0'),
 ('Friction between washers', 'text'),
 ('Unit price ($ per washer)', 'price'),
 ('Washer cost ($)', 'money'),
]
start = 5
for k,(lab,fmt) in enumerate(labels):
    r = start + k
    ws.cell(r,1,lab).font = black
    ws.cell(r,1).border = box
for col, wr, npar, fric in (('B',3,1,'None (series only)'),('C',4,2,'Yes, at every pair')):
    wrow = wr  # washer data row
    R = {lab:start+k for k,(lab,_) in enumerate(labels)}
    def c(lab): return f'{col}{R[lab]}'
    f = {}
    f['Washer size'] = f'=Washers!A{wrow}'
    f['Washers per group (parallel)'] = npar
    f['Per-washer force at preload (kN)'] = f'=Inputs!$B$2/{c("Washers per group (parallel)")}'
    f['Preload deflection (× h0)'] = '=' + inv(c('Per-washer force at preload (kN)'), wrow)
    f['Travel per group to limit (mm)'] = f'=(Inputs!$B$5-{c("Preload deflection (× h0)")})*Washers!E{wrow}'
    f['Minimum groups to stay within limit'] = f'=ROUNDUP(Inputs!$B$3/{c("Travel per group to limit (mm)")},0)'
    f['Groups in series (chosen)'] = 16
    f['Within deflection limit?'] = f'=IF({c("Release deflection (× h0)")}<=Inputs!$B$5,"Yes","No – add a group")'
    f['Total washers'] = f'={c("Groups in series (chosen)")}*{c("Washers per group (parallel)")}'
    f['Release deflection (× h0)'] = f'={c("Preload deflection (× h0)")}+Inputs!$B$3/{c("Groups in series (chosen)")}/Washers!E{wrow}'
    f['Release force (kN)'] = f'={c("Washers per group (parallel)")}*' + interp(c('Release deflection (× h0)'), wrow)
    f['Required at drawbar incl. P (kN)'] = f'={c("Release force (kN)")}+Inputs!$B$4'
    f['Cylinder margin'] = f'=Inputs!$B$10/{c("Required at drawbar incl. P (kN)")}-1'
    f['Group free height (mm)'] = f'=Washers!F{wrow}+({c("Washers per group (parallel)")}-1)*Washers!D{wrow}'
    f['Free stack length (mm)'] = f'={c("Groups in series (chosen)")}*{c("Group free height (mm)")}'
    f['Preloaded stack length (mm)'] = f'={c("Free stack length (mm)")}-{c("Groups in series (chosen)")}*{c("Preload deflection (× h0)")}*Washers!E{wrow}'
    f['Stack-height guide, 3 × OD (mm)'] = f'=3*Washers!B{wrow}'
    f['Stack-height check'] = f'=IF({c("Free stack length (mm)")}<={c("Stack-height guide, 3 × OD (mm)")},"OK","Exceeds")'
    f['Washer mass, total (g)'] = f'={c("Total washers")}*Washers!K{wrow}'
    f['Cap washer bore (mm)'] = f'=Washers!B{wrow}+Inputs!$B$12'
    f['Cap OD, estimate (mm)'] = f'={c("Cap washer bore (mm)")}+2*Inputs!$B$11'
    f['Friction between washers'] = fric
    f['Unit price ($ per washer)'] = None
    f['Washer cost ($)'] = f'=IF({c("Unit price ($ per washer)")}="","",{c("Unit price ($ per washer)")}*{c("Total washers")})'
    for k,(lab,fmt) in enumerate(labels):
        r = start+k
        cell = ws[f'{col}{r}']
        cell.value = f[lab]
        cell.border = box
        cell.font = blue if fmt=='in' else black
        if fmt=='price':
            cell.fill = yellow; cell.font = blue; cell.number_format = '$#,##0.00'
        elif fmt=='money': cell.number_format = '$#,##0.00'
        elif fmt not in ('text','in'): cell.number_format = fmt
        cell.alignment = Alignment(horizontal='right') if fmt!='text' else Alignment(horizontal='left')
ws.column_dimensions['A'].width = 36; ws.column_dimensions['B'].width = 20; ws.column_dimensions['C'].width = 20

n = start + len(labels) + 1
notes = [
 ('Why A50 single was favoured', True),
 ('1. Predictability: washers in a parallel pair rub as they flex. That friction adds hysteresis, so real release force is higher and clamp force lower and less repeatable than calculated, and preload set by stack height is less exact. Grease (MoS₂) reduces it but does not remove it. Series-only stacks avoid it.', False),
 ('2. Cylinder margin: A40 pairs leave less margin before friction. A 5.5:1 lever or running at 6.2 bar restores it.', False),
 ('3. Shorter stack and drawbar: about 23 mm less height. Headroom is not a constraint on this machine, so minor.', False),
 ('Where A40 pairs win', True),
 ('1. Smaller cap: turns from 70–75 mm bar instead of 85 mm, cheaper and less machining, with less rotating diameter at up to 2,500 rpm.', False),
 ('2. Availability: A40 is a common size and may be easier to get locally.', False),
 ('3. Cost, possibly: depends on the quote. 32 × A40 only wins if each is well under half the price of an A50; enter both unit prices above.', False),
 ('Bottom line', True),
 ('Both work; fatigue and washer mass are about equal. Let the quote decide: if A40s are cheap and in stock, use A40 pairs with a 5.5:1 lever and MoS₂ grease between washers; otherwise 16 × A50 single.', False),
 ('Method notes', True),
 ('Forces are linearly interpolated between the DIN 2093 tabulated points (0.25, 0.5, 0.75, 1.0 h0). Group count is an input (blue); the minimum count for the deflection limit is shown beside it. Results differ slightly from the chat figures (A50 10.4 kN, A40 11.2 kN), which assumed exactly 0.64 h0 at release rather than the deflection 16 groups actually gives. Friction is not modelled.', False),
]
for i,(t,isb) in enumerate(notes):
    r = n+i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r,1,t); c.font = bold if isb else black; c.alignment = wrap
    if not isb: ws.row_dimensions[r].height = 46
wb.save('PDB_washer_stack_comparison.xlsx')
